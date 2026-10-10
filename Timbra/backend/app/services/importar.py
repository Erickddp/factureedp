"""Importa clientes desde Excel (.xlsx), CSV o texto pegado desde Excel (tabuladores).

Encabezados reconocidos (en cualquier orden, con o sin acentos/guiones bajos):
  NOMBRE · RFC · CP · DIRECCION · REGIMEN_FISCAL · CORREO_O_CONTACTO     ← los de Rick
  USO_CFDI · TELEFONO · METODO_PAGO · FORMA_PAGO · NOTAS                  ← opcionales, van después
"""
import csv
import difflib
import io
import re

from satcfdi.catalogs import select_all

from .validacion import (GENERICOS, Resultado, quitar_acentos, usos_permitidos, validar_receptor_datos)

COLUMNAS = {
    "nombre": {"NOMBRE", "RAZON SOCIAL", "NOMBRE O RAZON SOCIAL", "CLIENTE"},
    "rfc": {"RFC"},
    "codigo_postal": {"CP", "C P", "CODIGO POSTAL", "CP FISCAL", "DOMICILIO FISCAL"},
    "direccion": {"DIRECCION", "DOMICILIO"},
    "regimen": {"REGIMEN FISCAL", "REGIMEN"},
    "contacto": {"CORREO O CONTACTO", "CORREO", "EMAIL", "E MAIL", "CONTACTO", "CORREOS"},
    "uso_cfdi": {"USO CFDI", "USO", "USO DEL CFDI"},
    "telefono": {"TELEFONO", "TEL", "CELULAR"},
    "metodo_pago": {"METODO PAGO", "METODO DE PAGO"},
    "forma_pago": {"FORMA PAGO", "FORMA DE PAGO"},
    "notas": {"NOTAS", "COMENTARIOS", "OBSERVACIONES"},
}

PLANTILLA_ENCABEZADOS = ["NOMBRE", "RFC", "CP", "DIRECCION", "REGIMEN_FISCAL", "CORREO_O_CONTACTO",
                         "USO_CFDI", "METODO_PAGO", "FORMA_PAGO", "NOTAS"]

ALIAS_REGIMEN = {
    "resico": "626", "simplificado de confianza": "626", "actividad empresarial": "612",
    "actividades empresariales": "612", "honorarios": "612", "arrendamiento": "606", "general de ley": "601",
    "sin obligaciones": "616", "plataformas": "625", "incorporacion fiscal": "621", "rif": "621",
    "asalariado": "605", "sueldos": "605", "fines no lucrativos": "603", "dividendos": "611",
}


def _norm(s: str) -> str:
    s = quitar_acentos(str(s or "")).upper()
    return re.sub(r"[^A-Z0-9]+", " ", s).strip()


def _norm_regimen(s: str) -> str:
    s = quitar_acentos(str(s or "")).lower()
    s = re.sub(r"^\s*(regimen\s+(de\s+(las\s+|los\s+)?)?)", "", s)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]+", " ", s)).strip()


_REGIMENES = {k: v for k, v in select_all("C756_c_RegimenFiscal").items()}
_REG_NORM = {_norm_regimen(v): k for k, v in _REGIMENES.items()}


def regimen_a_clave(texto: str) -> str | None:
    t = str(texto or "").strip()
    if not t:
        return None
    m = re.match(r"^\s*(\d{3})\b", t)
    if m and m.group(1) in _REGIMENES:
        return m.group(1)
    n = _norm_regimen(t)
    if n in _REG_NORM:
        return _REG_NORM[n]
    for alias, clave in ALIAS_REGIMEN.items():
        if alias in n:
            return clave
    cerca = difflib.get_close_matches(n, list(_REG_NORM), n=1, cutoff=0.72)
    return _REG_NORM[cerca[0]] if cerca else None


def _separar_contacto(texto: str) -> tuple[list[str], str | None]:
    emails, tels = [], []
    for tok in re.split(r"[|,;/\s]+", str(texto or "")):
        tok = tok.strip()
        if "@" in tok:
            emails.append(tok.lower())
        else:
            dig = re.sub(r"\D", "", tok)
            if len(dig) >= 10:
                tels.append(dig[-10:])
    return emails, (", ".join(tels) or None)


def leer_tabla(nombre_archivo: str | None, contenido: bytes | None, texto: str | None) -> list[list[str]]:
    if contenido and (nombre_archivo or "").lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        filas = []
        for row in ws.iter_rows(values_only=True):
            filas.append(["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
                          for v in row])
        return filas
    if contenido:
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                texto = contenido.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    texto = (texto or "").strip("\ufeff\r\n")
    if not texto:
        return []
    primera = texto.splitlines()[0]
    delim = "\t" if "\t" in primera else (";" if primera.count(";") > primera.count(",") else ",")
    return [row for row in csv.reader(io.StringIO(texto), delimiter=delim)]


def mapear_encabezados(encabezados: list[str]) -> dict[int, str]:
    mapa = {}
    for i, h in enumerate(encabezados):
        n = _norm(h)
        for campo, alias in COLUMNAS.items():
            if n in alias and campo not in mapa.values():
                mapa[i] = campo
                break
    return mapa


def uso_por_defecto(regimen: str | None, preferido: str) -> str:
    """Uso para filas sin USO_CFDI: el preferido si el régimen lo permite; si no G03 → S01 (siempre válido).
    D01 (honorarios médicos) sólo si el usuario lo elige: es correcto para pacientes, no para cualquiera."""
    permitidos = usos_permitidos(regimen or "")
    if preferido != "auto" and preferido in permitidos:
        return preferido
    for u in ("G03", "S01"):
        if u in permitidos:
            return u
    return "S01"


def previsualizar(filas: list[list[str]], existentes: dict[str, int], uso_default: str = "auto",
                  actuales: dict[str, list[dict]] | None = None, nombres: dict[str, set[str]] | None = None) -> dict:
    """`actuales` (rfc → registros guardados) activa la detección de conflictos: si el RFC ya existe con nombre, CP o régimen
    distintos, la fila queda "conflicto" y se decide fila por fila; si es idéntico y no trae nada nuevo, "igual"."""
    if not filas:
        return {"filas": [], "columnas": {}, "faltan": ["NOMBRE", "RFC", "CP", "REGIMEN_FISCAL"], "resumen": {}}
    mapa = mapear_encabezados(filas[0])
    faltan = [c.upper() for c in ("nombre", "rfc", "codigo_postal", "regimen") if c not in mapa.values()]
    out, vistos = [], {}
    for n, row in enumerate(filas[1:], start=2):
        if not any(str(x).strip() for x in row):
            continue
        crudo = {campo: (row[i] if i < len(row) else "") for i, campo in mapa.items()}
        rfc = re.sub(r"[\s\-]", "", str(crudo.get("rfc", ""))).upper()
        cp = re.sub(r"\D", "", str(crudo.get("codigo_postal", "")))
        cp = cp.zfill(5) if cp and len(cp) <= 5 else cp  # Excel se come el 0 inicial
        nombre = re.sub(r"\s+", " ", str(crudo.get("nombre", ""))).strip().upper()
        regimen_txt = str(crudo.get("regimen", "")).strip()
        regimen = regimen_a_clave(regimen_txt)
        emails, tel = _separar_contacto(crudo.get("contacto", ""))
        tel2 = re.sub(r"\D", "", str(crudo.get("telefono", "")))
        uso_txt = str(crudo.get("uso_cfdi", "")).strip().upper()[:4]
        uso = uso_txt or uso_por_defecto(regimen, uso_default)
        datos = {
            "rfc": rfc, "nombre": nombre, "codigo_postal": cp, "regimen_fiscal": regimen or "",
            "uso_cfdi": uso, "direccion": str(crudo.get("direccion", "")).strip() or None,
            "emails": emails, "telefono": tel2[-10:] if len(tel2) >= 10 else tel,
            "metodo_pago": (str(crudo.get("metodo_pago", "")).strip().upper()[:3] or None),
            "forma_pago": (re.sub(r"\D", "", str(crudo.get("forma_pago", "")))[:2].zfill(2)
                           if str(crudo.get("forma_pago", "")).strip() else None),
            "notas": str(crudo.get("notas", "")).strip() or None,
        }
        r = Resultado()
        if regimen is None:
            r.error("regimen_fiscal", f"No reconozco el régimen «{regimen_txt or '(vacío)'}». Pon la clave (ej. 605) o el nombre del catálogo.")
        validar_receptor_datos(r, rfc, nombre, regimen or "", cp, uso)
        if datos["metodo_pago"] and datos["metodo_pago"] not in ("PUE", "PPD"):
            r.error("metodo_pago", "Método de pago debe ser PUE o PPD.")
        if rfc in GENERICOS:
            r.aviso("rfc", "RFC genérico: no hace falta darlo de alta como cliente.")
        # quitar avisos del nombre si sólo es el societario y ofrecer la versión limpia
        for a in r.avisos:
            if a["campo"].endswith("nombre") and a.get("sugerencia"):
                datos["nombre"] = a["sugerencia"]
                a["mensaje"] = "Le quité el régimen societario al nombre (CFDI 4.0 va sin él)."
        accion = "error" if r.errores else ("actualizar" if rfc in existentes else "nuevo")
        extra: dict = {}
        if actuales is not None and not r.errores:
            regs = actuales.get(rfc) or []
            if regs:
                base = next((x for x in regs if _norm(x["nombre"]) == _norm(nombre)), regs[0])
                difs = []
                for campo, etiqueta in (("nombre", "Nombre"), ("codigo_postal", "CP"), ("regimen_fiscal", "Régimen")):
                    h, b = datos[campo], base[campo]
                    distinto = _norm(h) != _norm(b) if campo == "nombre" else h != b
                    if h and distinto:
                        difs.append({"campo": campo, "etiqueta": etiqueta, "hoja": h, "base": b})
                trae_nuevo = (bool(set(emails) - set(base["emails"])) or (bool(datos["telefono"]) and not base["telefono"])
                              or (bool(datos["direccion"]) and not base["direccion"]))
                accion = "conflicto" if difs else ("actualizar" if trae_nuevo else "igual")
                extra = {"cliente_id": base["id"], "conflictos": difs}
                if len(regs) > 1:
                    r.aviso("rfc", f"Ya hay {len(regs)} registros guardados con este RFC (revísalos en Clientes → Revisar → Duplicados).")
            elif nombres and _norm(nombre) in nombres and nombres[_norm(nombre)] - {rfc}:
                otros = ", ".join(sorted(nombres[_norm(nombre)] - {rfc}))
                r.aviso("nombre", f"Ya tienes un cliente con este nombre pero otro RFC ({otros}). ¿Es la misma persona?")
        if rfc in vistos and not r.errores:
            out[vistos[rfc]]["accion"] = "duplicado"
            out[vistos[rfc]]["avisos"].append({"campo": "rfc", "mensaje": f"Se repite en la fila {n}; se usa la última."})
        vistos[rfc] = len(out)
        out.append({"fila": n, "datos": datos, "errores": r.errores, "avisos": r.avisos, "accion": accion,
                    "regimen_original": regimen_txt, **extra})
    resumen = {k: sum(1 for f in out if f["accion"] == k)
               for k in ("nuevo", "actualizar", "error", "duplicado", "conflicto", "igual")}
    return {"filas": out, "columnas": {str(k): v for k, v in mapa.items()}, "faltan": faltan, "resumen": resumen}


def plantilla_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append(PLANTILLA_ENCABEZADOS)
    ws.append(["MARIA EJEMPLO PEREZ", "XAXX010101000", "06600", "CALLE 1 #2, COL. CENTRO, CDMX",
               "Sueldos y Salarios e Ingresos Asimilados a Salarios", "correo@ejemplo.com | 5512345678",
               "D01", "PUE", "03", "Borra esta fila de ejemplo"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="152033")
    for col, w in zip("ABCDEFGHIJ", (38, 16, 8, 45, 48, 40, 10, 12, 12, 30)):
        ws.column_dimensions[col].width = w
    for fila in ws.iter_rows(min_row=2, max_row=500, min_col=3, max_col=3):
        for c in fila:
            c.number_format = "@"  # CP como texto: no pierde el 0 inicial
    cat = wb.create_sheet("Regimenes")
    cat.append(["CLAVE", "REGIMEN FISCAL", "USOS CFDI PERMITIDOS"])
    for k, v in sorted(_REGIMENES.items()):
        cat.append([k, v, ", ".join(usos_permitidos(k))])
    cat.column_dimensions["B"].width = 70
    cat.column_dimensions["C"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
