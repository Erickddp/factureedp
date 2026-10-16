"""Lo que tarda en el módulo de XML: bajar del SAT, pintar PDF, armar el ZIP y el Excel.

La descarga usa el flujo de la app de escritorio (xmlsat/flujo.py), que ya está probado contra el SAT real. Lo añadido
aquí son los tramos «meta_*»: la lista (Metadata) con la que el SAT dice qué comprobantes están cancelados.
"""
import base64
import io
import json
import re
import tempfile
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable

from sqlalchemy.orm import Session

from ..db import SessionLocal
from ..models import Emisor, XmlCfdi
from ..xmlsat import pdf as pdf_nucleo
from ..xmlsat.config import Opciones, Perfil
from ..xmlsat.flujo import COD_OK, Descargador, Resumen
from ..xmlsat.lector import leer_detalle
from ..xmlsat.utilerias import con_reintentos
from . import xml_firma, xml_indice
from .xml_trabajos import Trabajo

# Las pruebas ponen aquí un SAT simulado: fabrica_cliente(emisor_id) -> objeto con token/solicitar/verificar/descargar_paquete.
fabrica_cliente: Callable | None = None
POLL_SEG = 30          # el flujo espera 30 s, luego 60 y 120 entre verificaciones
POLL_INTENTOS = 40
LIMITE_ESPERA_SEG = 150 * 60  # cuánto se espera al SAT en una corrida (la e.firma vive 3 h en memoria)
REVISAR_ESTATUS_CADA = timedelta(hours=24)  # la lista de cancelados de un periodo se vuelve a pedir pasado este tiempo
PIE_PDF = "Generado con Timbra"


class SinEfirma(RuntimeError):
    pass


def _cliente(emisor_id: int):
    if fabrica_cliente is not None:
        return fabrica_cliente(emisor_id)
    fiel = xml_firma.tomar(emisor_id)
    if fiel is None:
        raise SinEfirma("Pon tu e.firma para hablar con el SAT. Por seguridad no se guarda: se pide cada vez.")
    from cfdiclient import SolicitaDescargaEmitidos, SolicitaDescargaRecibidos

    from ..xmlsat.cliente import ClienteSAT

    class _Real(ClienteSAT):
        def solicitar_metadata(self, rfc, tipo, desde, hasta):
            """La lista de comprobantes del periodo con su estatus (vigentes Y cancelados): por eso no se filtra por estado."""
            if tipo == "emitidos":
                return SolicitaDescargaEmitidos(self.fiel, timeout=self.timeout).solicitar_descarga(
                    self.token(), rfc, desde, hasta, rfc_emisor=rfc, tipo_solicitud="Metadata", estado_comprobante=None)
            return SolicitaDescargaRecibidos(self.fiel, timeout=self.timeout).solicitar_descarga(
                self.token(), rfc, desde, hasta, rfc_receptor=rfc, tipo_solicitud="Metadata", estado_comprobante=None)

    return _Real(fiel)


# Respuestas con las que el SAT rechaza a quien firma: reintentar o seguir con los demás meses no sirve de nada.
RECHAZOS_FIRMA = {
    "300": "El SAT no reconoce al usuario de esa e.firma.",
    "302": "El SAT no pudo leer la firma de la solicitud.",
    "303": "La e.firma no corresponde al RFC que pide la descarga.",
    "304": "El SAT dice que esa e.firma está revocada o caduca. Renuévala y pon la nueva.",
    "305": "El SAT no acepta ese certificado. Suele pasar con e.firmas de prueba (como la del RFC de pruebas EKU9003173C9), "
           "con un CSD en lugar de la e.firma, o con una e.firma que ya fue sustituida: usa la e.firma vigente del contribuyente.",
}


class _ConMetadata:
    """Para el flujo, «meta_emitidos» es un tipo más; aquí se traduce a una solicitud de Metadata."""

    def __init__(self, cli, al_rechazo: Callable[[str], None] = lambda m: None):
        self.cli = cli
        self.al_rechazo = al_rechazo
        self.rechazo: str | None = None

    def token(self):
        return self.cli.token()

    def solicitar(self, rfc, tipo, desde, hasta):
        if tipo.startswith("meta_"):
            r = self.cli.solicitar_metadata(rfc, tipo[5:], desde, hasta)
        else:
            r = self.cli.solicitar(rfc, tipo, desde, hasta)
        cod = str((r or {}).get("cod_estatus") or "")
        if cod in RECHAZOS_FIRMA and self.rechazo is None:
            self.rechazo = f"{RECHAZOS_FIRMA[cod]} (respuesta del SAT: {cod} {(r or {}).get('mensaje') or ''})".strip()
            self.al_rechazo(self.rechazo)
        return r

    def verificar(self, rfc, id_solicitud):
        v = self.cli.verificar(rfc, id_solicitud)
        # Visto con el SAT real: cuando en el periodo no hay comprobantes, la solicitud termina «rechazada» (estado 5) y el
        # motivo viene en CodigoEstadoSolicitud = 5004 (sin información). No es un error: el periodo está vacío.
        if isinstance(v, dict) and str(v.get("codigo_estado_solicitud") or "") == "5004":
            v = {**v, "cod_estatus": "5004"}
        return v

    def descargar_paquete(self, rfc, id_paquete):
        return self.cli.descargar_paquete(rfc, id_paquete)


class _Descargador(Descargador):
    """El flujo de escritorio tal cual; sólo cambia qué se hace con el paquete cuando es una lista de estatus."""

    al_metadata: Callable[[str, str], None] = staticmethod(lambda tipo, texto: None)

    def _descargar_y_guardar(self, tipo, id_paquete, etiqueta, res):
        if not tipo.startswith("meta_"):
            return super()._descargar_y_guardar(tipo, id_paquete, etiqueta, res)
        d = con_reintentos(lambda: self.cli.descargar_paquete(self.rfc, id_paquete), self.op.reintentos,
                           self.op.backoff_base_seg, self.log, f"{etiqueta} descargar {id_paquete}", self.dormir)
        if d.get("cod_estatus") not in (COD_OK, None) or not d.get("paquete_b64"):
            raise RuntimeError(f"lista {id_paquete} no descargó: cod={d.get('cod_estatus')} {d.get('mensaje')}")
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(d["paquete_b64"]))) as z:
            for info in z.infolist():
                if not info.is_dir() and info.filename.lower().endswith(".txt"):
                    self.al_metadata(tipo[5:], z.read(info).decode("utf-8", "replace"))
        res.paquetes += 1


def probar_efirma(emisor_id: int) -> tuple[bool, str]:
    """Pide un token al SAT: si lo da, la e.firma sirve para descargar."""
    try:
        cliente = _cliente(emisor_id)
    except SinEfirma as e:
        return False, str(e)
    try:
        cliente.token()
        return True, "El SAT te dio acceso con esa e.firma. Ya puedes descargar."
    except Exception as e:  # noqa: BLE001
        return False, f"La e.firma abrió, pero el SAT no respondió bien: {str(e)[:300]}"


def _refrescar_estatus(archivo_estado: str, ahora: datetime) -> None:
    """Las listas de estatus (meta_*) ya completadas se vuelven a pedir pasado un día: un CFDI se puede cancelar después.
    Los XML no: el SAT limita cuántas veces se pide el mismo periodo, y un XML ya bajado no cambia."""
    p = Path(archivo_estado)
    try:
        datos = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return
    cambio = False
    for k in list(datos):
        v = datos[k]
        if not k.startswith("meta_") or not v.get("completado"):
            continue
        try:
            viejo = ahora - datetime.fromisoformat(v.get("actualizado", "")) > REVISAR_ESTATUS_CADA
        except ValueError:
            viejo = True
        if viejo:
            datos.pop(k)
            cambio = True
    if cambio:
        p.write_text(json.dumps(datos, indent=2, ensure_ascii=False), encoding="utf-8")


def descargar_sat(t: Trabajo, emisor_id: int, rfc: str, nombre: str, desde: datetime, hasta: datetime, tipo: str,
                  con_estatus: bool = True) -> dict:
    """Solicita, espera y baja los paquetes del SAT; cada XML nuevo entra al índice. Después (si se pide) baja la lista
    de estatus para marcar los cancelados. Un tropiezo con la lista no echa a perder los XML ya bajados."""
    bit = t.bitacora()
    try:
        cliente = _ConMetadata(_cliente(emisor_id), al_rechazo=lambda m: t.cancelar.set())  # no seguir con los demás meses
    except SinEfirma as e:
        return {"errores": [str(e)]}
    perfil = Perfil(rfc=rfc, nombre=nombre or rfc, carpeta_estado=str(xml_indice.carpeta_estado()), id=str(emisor_id))
    op = Opciones(carpeta_salida=str(xml_indice.carpeta_xml(emisor_id)), poll_intervalo_seg=POLL_SEG, poll_max_intentos=POLL_INTENTOS)
    _refrescar_estatus(perfil.archivo_estado(op), datetime.now())
    nuevos = [0]
    estatus = {"revisados": 0, "cancelados": 0, "nuevos_cancelados": 0, "sin_xml": 0}

    def al_guardar(ruta: Path) -> None:
        with SessionLocal() as db:
            xml_indice.indexar(db, emisor_id, ruta)
            db.commit()
        nuevos[0] += 1
        if nuevos[0] % 25 == 0:
            t.avance(nuevos[0], 0, f"{nuevos[0]} XML nuevos")

    def al_metadata(tipo_meta: str, texto: str) -> None:
        with SessionLocal() as db:
            r = xml_indice.aplicar_metadata(db, db.get(Emisor, emisor_id), "Emitido" if tipo_meta == "emitidos" else "Recibido", texto)
        for k in estatus:
            estatus[k] += r[k]

    d = _Descargador(perfil, op, cliente, bit, dormir=lambda s: t.cancelar.wait(s), clasificar=al_guardar, cancelado=t.cancelar.is_set)
    d.al_metadata = al_metadata
    tipos = ["emitidos", "recibidos"] if tipo == "ambos" else [tipo]
    todos = tipos + ([f"meta_{x}" for x in tipos] if con_estatus else [])
    t.avance(0, 0, "Pidiendo al SAT todos los meses a la vez…")

    def avance(listos: int, total_pedidos: int) -> None:
        t.avance(nuevos[0], 0, f"El SAT tiene listos {listos} de {total_pedidos} pedidos · {nuevos[0]} XML nuevos"
                 if total_pedidos else f"{nuevos[0]} XML nuevos")
    # todos los meses y tipos se piden de una vez y se bajan conforme el SAT los tiene listos (antes era uno por uno)
    total = d.ejecutar_paralelo(desde, hasta, todos, limite_seg=LIMITE_ESPERA_SEG, avance=avance)
    avisos: list[str] = [e for e in total.errores if e.startswith("meta_")]
    total.errores = [e for e in total.errores if not e.startswith("meta_")]
    if total.pendientes:
        avisos.insert(0, f"El SAT sigue preparando {len(total.pendientes)} pedido(s) ({', '.join(total.pendientes[:6])}"
                         f"{'…' if len(total.pendientes) > 6 else ''}). Ya quedaron pedidos: vuelve en un rato, pon tu e.firma y pide "
                         "el mismo periodo; se bajan sin volver a pedirlos.")
    if cliente.rechazo:  # un solo mensaje claro, no uno por mes; y la e.firma que no sirve se olvida
        t.cancelar.clear()
        xml_firma.olvidar(emisor_id)
        return {"tramos": total.tramos, "solicitudes": total.solicitudes_nuevas, "paquetes": total.paquetes, "nuevos": total.xml_nuevos,
                "repetidos": total.xml_duplicados, "errores": [cliente.rechazo], "avisos": [], "estatus": None}
    return {"tramos": total.tramos, "solicitudes": total.solicitudes_nuevas, "paquetes": total.paquetes, "nuevos": total.xml_nuevos,
            "repetidos": total.xml_duplicados, "errores": total.errores, "avisos": avisos[:20], "pendientes": total.pendientes,
            "estatus": estatus if con_estatus else None}


# ---------------------------------------------------------------- PDF propio (una hoja, sin navegador)
def diseno_de(emisor: Emisor) -> dict:
    """El PDF lleva el nombre del contribuyente, su color y su logo (los mismos de sus facturas)."""
    d = {"titulo": "", "subtitulo": "Representación impresa de CFDI", "pie": PIE_PDF, "logo": ""}
    color = (emisor.color or "").strip()
    if re.fullmatch(r"#[0-9a-fA-F]{6}", color):
        d["color_primario"] = color
    m = re.match(r"data:image/(png|jpe?g);base64,(.+)", emisor.logo_b64 or "", re.S)
    if m:
        try:
            ruta = xml_indice.carpeta_tmp() / f"logo_{emisor.id}.{'png' if m.group(1) == 'png' else 'jpg'}"
            ruta.write_bytes(base64.b64decode(m.group(2)))
            d["logo"] = str(ruta)
        except (ValueError, OSError):
            pass
    return d


def pintar_pdf(emisor_id: int, diseno: dict, d: dict, destino: Path, resumen: dict | None = None) -> None:
    """d: primer renglón del comprobante (ya con seguimiento de pago si es factura)."""
    if not d.get("Archivo"):
        raise ValueError("de este comprobante sólo se sabe que existe (lo avisó el SAT): no hay XML para pintarlo")
    det = leer_detalle(xml_indice.carpeta_xml(emisor_id) / d["Archivo"])
    if not det:
        raise ValueError("el XML ya no está o no es un CFDI legible")
    pdf_nucleo.cfdi_a_pdf(det, str(destino), config=diseno, clase=d.get("Clase", ""), tipo=d.get("Tipo", ""),
                          inteligencia=xml_indice.seguimiento(d, resumen or {}))


def nombre_pdf(d: dict) -> str:
    return pdf_nucleo.nombre_pdf(pdf_nucleo.CONFIG_DEFAULT["general"]["nombre_archivo"], d)


def pdf_de(db: Session, emisor: Emisor, x: XmlCfdi) -> tuple[str, bytes]:
    """(nombre, contenido) del PDF de un comprobante."""
    from ..xmlsat.lector import Renglon
    filas = [Renglon(dict(r)) for r in x.filas]
    resumen = xml_indice.con_inteligencia(db, emisor.id, filas) if x.comprobante == "Ingreso" else {}
    d = filas[0].datos
    with tempfile.TemporaryDirectory(dir=xml_indice.carpeta_tmp()) as tmp:
        salida = Path(tmp) / "cfdi.pdf"
        pintar_pdf(emisor.id, diseno_de(emisor), d, salida, resumen)
        return nombre_pdf(d), salida.read_bytes()


# ---------------------------------------------------------------- ZIP
def armar_zip(t: Trabajo, emisor_id: int, f: xml_indice.Filtros, incluir: str) -> dict:
    """ZIP con los XML, los PDF o ambos, en las mismas carpetas de la app de escritorio:
    {AAAAMM} {4 letras del RFC}/{emitidos|recibidos}/. El XML se llama como su UUID; el PDF, con nombre legible."""
    con_xml, con_pdf = incluir in ("xml", "ambos"), incluir in ("pdf", "ambos")
    with SessionLocal() as db:
        emisor = db.get(Emisor, emisor_id)
        diseno = diseno_de(emisor)
        filas = xml_indice.renglones(db, f)
        resumen = xml_indice.con_inteligencia(db, emisor_id, filas) if con_pdf else {}
    unicos, vistos = [], set()
    for r in filas:  # un pago con tres documentos son tres renglones y un solo PDF
        arch = r.datos.get("Archivo")
        if arch and arch not in vistos:
            vistos.add(arch)
            unicos.append(r.datos)
    base = xml_indice.carpeta_xml(emisor_id)
    destino = xml_indice.carpeta_tmp() / f"paquete_{t.id}.zip"
    errores, usados = [], set()
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z, tempfile.TemporaryDirectory(dir=xml_indice.carpeta_tmp()) as tmp:
        for i, d in enumerate(unicos, 1):
            if t.cancelar.is_set():
                break
            carpeta = str(Path(d["Archivo"]).parent.as_posix())
            try:
                if con_xml:
                    z.write(base / d["Archivo"], d["Archivo"])
                if con_pdf:
                    nombre = f"{carpeta}/{nombre_pdf(d)}"
                    if nombre in usados:  # mismo folio emitido dos veces: que no se pisen
                        nombre = f"{nombre[:-4]}_{d['UUID'][-12:]}.pdf"
                    usados.add(nombre)
                    salida = Path(tmp) / "cfdi.pdf"
                    pintar_pdf(emisor_id, diseno, d, salida, resumen)
                    z.write(salida, nombre)
            except Exception as e:  # noqa: BLE001 — un XML malo no detiene el paquete
                errores.append(f"{d.get('UUID') or d['Archivo']}: {e}")
            if i % 10 == 0 or i == len(unicos):
                t.avance(i, len(unicos), f"{i} de {len(unicos)}")
    if t.cancelar.is_set():
        destino.unlink(missing_ok=True)
        return {"comprobantes": 0}
    return {"comprobantes": len(unicos), "avisos": errores[:50], "archivo": destino.name}


# ---------------------------------------------------------------- Excel
def excel(db: Session, emisor: Emisor, f: xml_indice.Filtros) -> Path:
    """Hoja «CFDI»: EXACTAMENTE las 65 columnas del formato maestro, en su orden (es la que se pega en la hoja del despacho).
    Lo que no es del formato va aparte: «Seguimiento» (clase, pagos, estatus del SAT, archivo), «Resumen» y «Nómina»."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from ..xmlsat.exportar import _hoja_resumen
    from ..xmlsat.lector import NUMERICAS

    filas = xml_indice.renglones(db, f)
    xml_indice.con_inteligencia(db, emisor.id, filas)
    estatus = {(u, tp): (e, c or "") for u, tp, e, c in xml_indice.consulta(db, f).with_entities(
        XmlCfdi.uuid, XmlCfdi.tipo, XmlCfdi.estatus, XmlCfdi.cancelado_el)}
    destino = xml_indice.carpeta_tmp() / f"excel_{emisor.id}_{id(filas)}.xlsx"

    def hoja(ws, titulos: list[str], renglones: list[list], numericas=frozenset()):
        ws.append(titulos)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center", vertical="center")
        for r in renglones:
            ws.append(r)
        for i, col in enumerate(titulos, start=1):
            letra = get_column_letter(i)
            if col in numericas:
                fmt = "#,##0" if col == "Conceptos" else "#,##0.00"
                for celda in ws[letra][1:]:
                    celda.number_format = fmt
            largos = [len(col)] + [len(str(r[i - 1] if r[i - 1] is not None else "")) for r in renglones[:400]]
            ws.column_dimensions[letra].width = min(max(largos) + 2, 42)
        ws.freeze_panes = "A2"
        if renglones:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{len(renglones) + 1}"

    wb = Workbook()
    maestro = xml_indice.FORMATO_EDDP
    ws = wb.active
    ws.title = "CFDI"
    hoja(ws, maestro, [[r.datos.get(c, "") for c in maestro] for r in filas], NUMERICAS)

    extra = ["UUID", "Tipo", "Serie", "Folio", "Clase", "Estatus Pago", "Pagado REP", "Num REP", "NC Aplicadas", "Saldo Pendiente", "Factura Encontrada",
             "Estatus SAT", "Cancelado el", "Archivo"]
    seg = []
    for r in filas:
        e, c = estatus.get(((r.datos.get("UUID") or "").upper(), r.datos.get("Tipo")), ("Vigente", ""))
        seg.append([{"Estatus SAT": e, "Cancelado el": c}.get(k, r.datos.get(k, "")) for k in extra])
    hoja(wb.create_sheet("Seguimiento"), extra, seg, {"Pagado REP", "NC Aplicadas", "Saldo Pendiente"})
    _hoja_resumen(wb, filas)

    recibos = [r.datos for r in filas if r.datos.get("Clase") == "N" and r.datos.get("_nomina")]
    if recibos:
        titulos = ["Tipo", "Fecha pago", "Periodo desde", "Periodo hasta", "Días", "RFC patrón", "Patrón", "RFC trabajador", "Trabajador", "Puesto",
                   "Percepciones", "Gravado", "Exento", "Otros pagos", "Subsidio", "Deducciones", "ISR retenido", "Seguridad social", "Otras deducciones",
                   "Neto pagado", "Serie", "Folio", "UUID"]
        datos = []
        for d in sorted(recibos, key=lambda x: x["_nomina"].get("fecha_pago") or x.get("Fecha") or ""):
            n = d["_nomina"]
            datos.append([d.get("Tipo"), n.get("fecha_pago"), n.get("desde"), n.get("hasta"), n.get("dias"), d.get("RFC Emisor"), d.get("Nombre Emisor"),
                          d.get("RFC Receptor"), d.get("Nombre Receptor"), n.get("puesto"), n.get("percepciones"), n.get("gravado"), n.get("exento"),
                          n.get("otros_pagos"), n.get("subsidio"), n.get("deducciones"), n.get("isr"), n.get("imss"), n.get("otras_deducciones"),
                          d.get("Total"), d.get("Serie"), d.get("Folio"), d.get("UUID")])
        hoja(wb.create_sheet("Nómina"), titulos, datos, set(titulos[10:20]))
    wb.save(destino)
    return destino
