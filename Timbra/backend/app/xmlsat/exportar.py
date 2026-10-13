# -*- coding: utf-8 -*-
"""Exportación: Excel con todas las columnas y texto tabulado con las columnas pedidas
(el formato EDDP que se pega en la hoja maestra). Mismas funciones de la app de escritorio."""
from pathlib import Path
from typing import Dict, List, Tuple

from .lector import COLUMNAS, NUMERICAS, Renglon


# ---------------------------------------------------------------- Excel
def a_excel(filas: List[Renglon], destino: str, hoja: str = "CFDI") -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    d = Path(destino)
    d.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = hoja

    ws.append(COLUMNAS)
    encabezado = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = encabezado
        c.fill = relleno
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r in filas:
        ws.append(r.como_lista())

    # formato de número y ancho de columna
    for i, col in enumerate(COLUMNAS, start=1):
        letra = get_column_letter(i)
        if col in NUMERICAS:
            fmt = "#,##0" if col == "Conceptos" else "#,##0.00"
            for celda in ws[letra][1:]:
                celda.number_format = fmt
        largos = [len(col)] + [len(str(r.datos.get(col, ""))) for r in filas[:400]]
        ws.column_dimensions[letra].width = min(max(largos) + 2, 42)

    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{len(filas) + 1}"

    _hoja_resumen(wb, filas)
    wb.save(d)
    return str(d)


def _hoja_resumen(wb, filas: List[Renglon]) -> None:
    from openpyxl.styles import Font
    ws = wb.create_sheet("Resumen")
    ws.append(["Mes", "Tipo", "Comprobantes", "SubTotal", "IVA Trasladado",
               "IVA Retenido", "ISR Retenido", "Total"])
    for c in ws[1]:
        c.font = Font(bold=True)

    agrupado: Dict[Tuple[str, str], List[float]] = {}
    for r in filas:
        k = (f"{r.datos.get('Año','')}-{r.datos.get('Mes','')}", str(r.datos.get("Tipo", "")))
        a = agrupado.setdefault(k, [0, 0.0, 0.0, 0.0, 0.0, 0.0])
        a[0] += 1
        for i, col in enumerate(["SubTotal", "IVA Trasladado", "IVA Retenido",
                                 "ISR Retenido", "Total"], start=1):
            a[i] += float(r.datos.get(col) or 0)

    for (mes, tipo), a in sorted(agrupado.items()):
        ws.append([mes, tipo, a[0], round(a[1], 2), round(a[2], 2),
                   round(a[3], 2), round(a[4], 2), round(a[5], 2)])
    for fila in ws.iter_rows(min_row=2, min_col=4):
        for celda in fila:
            celda.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12


# ---------------------------------------------------------------- texto tabulado
def a_tsv_cols(filas, columnas, con_encabezado=True):
    """TSV solo con las columnas pedidas (para copiar/CSV configurable)."""
    def limpiar(v):
        s = "" if v is None else str(v)
        return s.replace("\t", " ").replace("\r", " ").replace("\n", " ")
    lineas = []
    if con_encabezado:
        lineas.append("\t".join(columnas))
    for r in filas:
        lineas.append("\t".join(limpiar(r.datos.get(c, "")) for c in columnas))
    return "\n".join(lineas)
